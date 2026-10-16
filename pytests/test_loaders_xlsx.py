"""Чтение XLSX: выбор листа, числа и даты Excel, координаты ячеек (буква колонки, строка)."""

from __future__ import annotations

import datetime as dt
from collections.abc import Sequence
from pathlib import Path

from openpyxl import Workbook

from iolistgen.config import InputConfig
from iolistgen.loaders import load_table
from iolistgen.validate import check_file
from pytests.helpers import codes, config_from

HEADER = ["Тег", "Описание", "Тип сигнала", "Тип устройства", "Номер устройства", "Роль", "Мин.", "Макс.", "Ед. изм."]


def book(path: Path, sheets: dict[str, Sequence[Sequence[object]]]) -> Path:
    workbook = Workbook()
    first = True
    for title, rows in sheets.items():
        sheet = workbook.active if first else workbook.create_sheet()
        assert sheet is not None
        sheet.title = title
        first = False
        for row in rows:
            sheet.append(list(row))
    workbook.save(path)
    return path


def test_numbers_and_header_detection(tmp_path: Path) -> None:
    path = book(
        tmp_path / "list.xlsx",
        {
            "Перечень": [
                ["Перечень сигналов КНС-1"],
                [],
                HEADER,
                ["LT101", "Уровень", "AI", "AnalogInput", 1, "PV", 0, 4.5, "м"],
                ["PT101", "Давление", "AI", "AnalogInput", 2.0, "PV", 0, 10, "бар"],
            ]
        },
    )
    table = load_table(path, InputConfig())
    assert table.format == "xlsx"
    assert table.sheet == "Перечень"
    assert table.header_row == 3
    result = check_file(path)
    assert result.ok, [d.render() for d in result.diagnostics]
    assert result.model is not None
    assert result.model.source.sheet == "Перечень"
    assert [d.dev_no for d in result.model.devices] == [1, 2]
    assert result.model.signals[0].range_max == 4.5


def test_first_sheet_with_header_is_chosen(tmp_path: Path) -> None:
    path = book(
        tmp_path / "list.xlsx",
        {
            "Титул": [["КНС-1"], ["Перечень сигналов"]],
            "Сигналы": [HEADER, ["LT101", "Уровень", "AI", "AnalogInput", 1, "PV", 0, 5, "м"]],
        },
    )
    result = check_file(path)
    assert result.ok
    assert result.source.sheet == "Сигналы"


def test_named_sheet_and_missing_sheet(tmp_path: Path) -> None:
    path = book(
        tmp_path / "list.xlsx",
        {
            "A": [HEADER, ["LT101", "Уровень", "AI", "AnalogInput", 1, "PV", 0, 5, "м"]],
            "B": [HEADER, ["LT201", "Уровень", "AI", "AnalogInput", 1, "PV", 5, 0, "м"]],
        },
    )
    assert check_file(path, sheet="a").ok
    result_b = check_file(path, config_from({"input": {"sheet": "B"}}))
    assert codes(result_b) == ["IOL204"]
    missing = check_file(path, sheet="C")
    assert codes(missing) == ["IOL001"]
    assert "«A»" in missing.diagnostics[0].hint


def test_cell_coordinates_in_messages(tmp_path: Path) -> None:
    path = book(
        tmp_path / "list.xlsx",
        {"Перечень": [HEADER, ["LT101", "Уровень", "AI", "AnalogInput", 1, "PV", 0, 5, None]]},
    )
    result = check_file(path)
    assert codes(result) == ["IOL206"]
    source = result.diagnostics[0].source
    assert source is not None
    assert (source.sheet, source.row, source.column_letter, source.header) == ("Перечень", 2, "I", "Ед. изм.")
    assert "лист «Перечень»" in result.diagnostics[0].render()


def test_excel_dates_and_booleans_are_not_numbers(tmp_path: Path) -> None:
    path = book(
        tmp_path / "list.xlsx",
        {"Перечень": [HEADER, ["LT101", "Уровень", "AI", "AnalogInput", 1, "PV", dt.datetime(2026, 5, 1), True, "м"]]},
    )
    result = check_file(path)
    assert codes(result) == ["IOL007", "IOL007"]
    assert "дату" in result.diagnostics[0].hint


def test_broken_workbook(tmp_path: Path) -> None:
    path = tmp_path / "list.xlsx"
    path.write_bytes(b"not a zip")
    result = check_file(path)
    assert codes(result) == ["IOL001"]


def test_workbook_without_header(tmp_path: Path) -> None:
    path = book(tmp_path / "list.xlsx", {"Лист1": [["a", "b"], [1, 2]]})
    assert codes(check_file(path)) == ["IOL002"]


def test_formula_without_cached_value(tmp_path: Path) -> None:
    path = book(
        tmp_path / "list.xlsx",
        {"Перечень": [HEADER, ["LT101", "Уровень", "AI", "AnalogInput", 1, "PV", 0, "=2+3", "м"]]},
    )
    result = check_file(path)
    formula = [d for d in result.diagnostics if d.code == "IOL007"]
    assert len(formula) == 1
    assert formula[0].source is not None and formula[0].source.column_letter == "H"
    assert "=2+3" in formula[0].message
    assert "IOL205" in codes(result)  # пустая граница диапазона — следствие, но причина названа


def _patch_sheet_xml(path: Path, replacements: dict[str, str]) -> None:
    """Подменить текст XML листа: openpyxl пишет числа короче, чем Excel (17 значащих цифр)."""
    import zipfile

    with zipfile.ZipFile(path) as source:
        items = {name: source.read(name) for name in source.namelist()}
    sheet = items["xl/worksheets/sheet1.xml"].decode("utf-8")
    for old, new in replacements.items():
        assert old in sheet
        sheet = sheet.replace(old, new)
    items["xl/worksheets/sheet1.xml"] = sheet.encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for name, data in items.items():
            target.writestr(name, data)


def test_float_noise_from_formulas_is_rounded_like_excel(tmp_path: Path) -> None:
    """Значение формулы Excel 0.1+0.2 хранится как 0.30000000000000004; Excel показывает 15 знаков."""
    path = book(
        tmp_path / "list.xlsx",
        {"Перечень": [[*HEADER, "H"], ["LT101", "Уровень", "AI", "AnalogInput", 1, "PV", 0, 0.25, "м", 0.125]]},
    )
    _patch_sheet_xml(path, {"<v>0.25</v>": "<v>0.30000000000000004</v>", "<v>0.125</v>": "<v>0.27000000000000002</v>"})
    result = check_file(path)
    assert result.ok, [d.render() for d in result.diagnostics]
    assert result.model is not None
    assert result.model.signals[0].range_max == 0.3
    assert result.model.devices[0].setpoints.h == 0.27
    assert "0.30000000000000004" not in result.model.devices[0].cfg.initializer()


def test_excel_time_cells_as_delays(tmp_path: Path) -> None:
    path = book(
        tmp_path / "list.xlsx",
        {
            "Перечень": [
                [*HEADER, "HH", "Задержка вкл., с", "Задержка откл., с"],
                [
                    "LT101",
                    "Уровень",
                    "AI",
                    "AnalogInput",
                    1,
                    "PV",
                    0,
                    5,
                    "м",
                    4,
                    dt.time(0, 0, 5),
                    dt.timedelta(seconds=90),
                ],
            ]
        },
    )
    result = check_file(path)
    assert result.ok, [d.render() for d in result.diagnostics]
    assert result.model is not None
    device = result.model.devices[0]
    assert (device.delay_on_ms, device.delay_off_ms) == (5000, 90000)


def merged_book(path: Path) -> Path:
    """Типичная книга перечня: титул, раздел-подзаголовок, объединённые «Устройство» и «Группа HMI»."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Перечень"
    header: list[object] = [
        "Тег",
        "Описание",
        "Тип сигнала",
        "Устройство",
        "Тип устройства",
        "Номер устройства",
        "Роль",
        "Группа HMI",
    ]
    rows: list[list[object]] = [
        ["Перечень сигналов КНС-1"],
        header,
        ["Насосы"],
        ["M1_RUN", "Насос: работа", "DO", "M1", "Motor", 11, "CMD_RUN", "Насосы"],
        ["M1_FB", "Насос: в работе", "DI", None, None, None, "RUN_FB", None],
        ["M1_FLT", "Насос: авария", "DI", None, None, None, "FAULT", None],
        ["Резервуар"],
        ["LS101", "Поплавок", "DI", None, "DigitalInput", 1, "IN", "Резервуар"],
    ]
    for row in rows:
        sheet.append(row)
    sheet.merge_cells("A1:H1")  # титул
    sheet.merge_cells("A3:H3")  # подзаголовок раздела
    sheet.merge_cells("D4:D6")  # устройство M1 на три строки
    sheet.merge_cells("H4:H6")  # группа HMI
    sheet.merge_cells("A7:H7")
    workbook.save(path)
    return path


def test_vertically_merged_cells_apply_to_every_row(tmp_path: Path) -> None:
    result = check_file(merged_book(tmp_path / "list.xlsx"))
    assert result.ok, [d.render() for d in result.diagnostics]
    assert result.model is not None
    motor = result.model.device("M1")
    assert [s.tag for s in motor.signals] == ["M1_RUN", "M1_FB", "M1_FLT"]
    assert {s.hmi_group for s in motor.signals} == {"Насосы"}
    # строки-подзаголовки (объединение через всю таблицу) — не сигналы
    assert [s.tag for s in result.model.signals] == ["M1_RUN", "M1_FB", "M1_FLT", "LS101"]


def test_merged_cell_source_points_to_own_row(tmp_path: Path) -> None:
    table = load_table(merged_book(tmp_path / "list.xlsx"), InputConfig())
    rows = {row.cells["tag"].text: row for row in table.rows}
    device = rows["M1_FLT"].cells["device"]
    assert device.text == "M1"
    assert (device.source.row, device.source.column_letter) == (6, "D")


def test_content_hash_ignores_workbook_metadata(tmp_path: Path) -> None:
    """Пересохранение книги без правок не меняет SHA-256 перечня — и заголовки сгенерированных файлов."""
    data: list[object] = ["LT101", "Уровень", "AI", "AnalogInput", 1, "PV", 0, 4.5, "м"]
    rows: list[Sequence[object]] = [HEADER, data]
    first = book(tmp_path / "a.xlsx", {"Перечень": rows})
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Перечень"
    for row in rows:
        sheet.append(list(row))
    workbook.properties.modified = dt.datetime(2020, 1, 1)
    workbook.properties.lastModifiedBy = "другой инженер"
    second = tmp_path / "b.xlsx"
    workbook.save(second)
    assert first.read_bytes() != second.read_bytes()
    assert load_table(first, InputConfig()).sha256 == load_table(second, InputConfig()).sha256
    changed = book(tmp_path / "c.xlsx", {"Перечень": [HEADER, [*data[:7], 5, "м"]]})
    assert load_table(changed, InputConfig()).sha256 != load_table(first, InputConfig()).sha256
