"""Чтение перечня из XLSX (openpyxl): лист по имени или первый лист со строкой заголовка.

Читаются вычисленные значения (``data_only``): если в ячейке формула, а файл сохранён не в
Excel, значения у формулы нет и ячейка окажется пустой.

Объединённые ячейки читаются так, как их видит человек:

- объединение по вертикали ниже заголовка («Устройство», «Группа HMI», «Модуль» на несколько
  строк) — значение верхней ячейки действует в каждой строке объединения (в первой его колонке);
- строка, где заполнено только объединение по горизонтали (подзаголовок раздела «Насосная»
  через всю таблицу), — не сигнал и пропускается, как пустая.
"""

from __future__ import annotations

import hashlib
import io
import warnings
import zipfile
from collections.abc import Iterator, Sequence
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from iolistgen.columns import header_index
from iolistgen.config import InputConfig
from iolistgen.diagnostics import DiagnosticBag
from iolistgen.loaders.table import RawTable, build_table, is_empty, looks_like_header
from iolistgen.model import SourceRef


@dataclass(frozen=True)
class _Merges:
    """Объединённые ячейки листа ниже строки заголовка: что подставить и какие строки пропустить."""

    fill: dict[int, dict[int, object]]  # строка → колонка → значение верхней ячейки объединения
    captions: dict[int, int]  # строка-подзаголовок → колонка её объединения


def _merges(sheet: Worksheet, header_row: int | None) -> _Merges:
    fill: dict[int, dict[int, object]] = {}
    captions: dict[int, int] = {}
    if header_row is None:
        return _Merges(fill, captions)
    for area in sheet.merged_cells.ranges:
        if area.min_row <= header_row:
            continue  # титул и заголовок не размножаются
        value = sheet.cell(row=area.min_row, column=area.min_col).value
        if area.max_row > area.min_row:
            for row_no in range(area.min_row + 1, area.max_row + 1):
                fill.setdefault(row_no, {})[area.min_col] = value
        elif area.max_col > area.min_col:
            captions[area.min_row] = area.min_col
    return _Merges(fill, captions)


def _raw_records(sheet: Worksheet) -> Iterator[tuple[int, list[object]]]:
    width = sheet.max_column
    for row_no, values in enumerate(
        sheet.iter_rows(min_row=1, min_col=1, max_col=width, values_only=True),
        start=1,
    ):
        yield row_no, list(values)


def _records(sheet: Worksheet, merges: _Merges | None = None) -> Iterator[tuple[int, Sequence[object]]]:
    for row_no, values in _raw_records(sheet):
        if merges is not None:
            for column, value in merges.fill.get(row_no, {}).items():
                if column <= len(values) and values[column - 1] is None:
                    values[column - 1] = value
            caption = merges.captions.get(row_no)
            if caption is not None:
                filled = [position for position, value in enumerate(values, start=1) if not is_empty(value)]
                if filled == [caption]:
                    values = [None] * len(values)  # подзаголовок раздела через всю таблицу
        yield row_no, values


def _header_row(sheet: Worksheet, config: InputConfig) -> int | None:
    index = header_index(config.column_aliases)
    scanned = 0
    for row_no, values in _raw_records(sheet):
        if all(is_empty(value) for value in values):
            continue
        if looks_like_header(values, index):
            return row_no
        scanned += 1
        if scanned >= config.header_search_rows:
            return None
    return None


def load_xlsx(path: Path, config: InputConfig, *, sheet: str | None = None, display: str | None = None) -> RawTable:
    """Прочитать XLSX-файл перечня."""
    data = path.read_bytes()
    name = display or str(path)
    sha256 = hashlib.sha256(data).hexdigest()
    bag = DiagnosticBag()
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError, ValueError, TypeError) as exc:
        bag.add(
            "IOL001",
            f"файл не читается как книга Excel: {type(exc).__name__}: {exc}",
            "сохраните перечень в формате .xlsx (Excel 2007 и новее); .xls не поддерживается",
            source=SourceRef(file=name),
        )
        return RawTable(file=name, sheet=None, format="xlsx", sha256=sha256, diagnostics=bag)
    wanted = sheet or config.sheet
    worksheets = [ws for ws in workbook.worksheets if isinstance(ws, Worksheet)]
    chosen: Worksheet | None = None
    if wanted:
        chosen = next((ws for ws in worksheets if ws.title.casefold() == wanted.casefold()), None)
        if chosen is None:
            titles = ", ".join(f"«{ws.title}»" for ws in worksheets)
            bag.add(
                "IOL001",
                f"в книге нет листа «{wanted}»",
                f"укажите один из листов: {titles}",
                source=SourceRef(file=name),
            )
            return RawTable(file=name, sheet=wanted, format="xlsx", sha256=sha256, diagnostics=bag)
    else:
        chosen = next((ws for ws in worksheets if _header_row(ws, config) is not None), None)
        if chosen is None and worksheets:
            chosen = worksheets[0]
    if chosen is None:
        bag.add("IOL001", "в книге нет листов с данными", "добавьте лист с перечнем", source=SourceRef(file=name))
        return RawTable(file=name, sheet=None, format="xlsx", sha256=sha256, diagnostics=bag)
    table = build_table(
        file=name,
        sheet=chosen.title,
        fmt="xlsx",
        sha256=sha256,
        records=_records(chosen, _merges(chosen, _header_row(chosen, config))),
        config=config,
        bag=bag,
    )
    _report_formulas_without_values(data, chosen, table)
    return table


def _report_formulas_without_values(data: bytes, sheet: Worksheet, table: RawTable) -> None:
    """Формула без сохранённого значения читается как пустая ячейка — объяснить почему.

    Так бывает с книгами, которые записаны программой (openpyxl и т. п.) и не открывались в Excel:
    значения формул в файле нет, а пересчитывать формулы iolistgen не умеет.
    """
    if table.header_row is None:
        return
    empty = {
        (row_no, col_no)
        for row_no, values in _raw_records(sheet)
        if row_no > table.header_row
        for col_no, value in enumerate(values, start=1)
        if value is None
    }
    if not empty:
        return
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        formulas = load_workbook(io.BytesIO(data), data_only=False, read_only=False)[sheet.title]
    headers = {binding.index: binding.header or None for binding in table.columns}
    for row_no, col_no in sorted(empty):
        cell = formulas.cell(row=row_no, column=col_no)
        if cell.data_type == "f":
            table.diagnostics.add(
                "IOL007",
                f"в ячейке формула {cell.value} без сохранённого значения — ячейка прочитана как пустая",
                "откройте книгу в Excel и сохраните её: iolistgen читает сохранённые значения формул",
                source=SourceRef(
                    file=table.file, sheet=table.sheet, row=row_no, col=col_no, header=headers.get(col_no)
                ),
            )
