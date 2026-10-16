"""Перечень сигналов CSV → книга XLSX, как её ведут в Excel: титул, русские заголовки, числа.

Пример КНС хранится в CSV (его удобно ревьюить построчно), а XLSX собирается из него этим
скриптом: проверка XLSX тем же валидатором показывает, что оба формата дают одну модель.

Что делает:

- над таблицей — две строки титула (их загрузчик iolistgen пропускает: заголовок ищется по
  колонке «Тег»);
- заголовки колонок — русские названия из схемы (``iolistgen.columns``): «Тег», «Тип сигнала»,
  «Задержка вкл., с»; с ``--headers en`` — канонические английские имена;
- числовые колонки (диапазон, уставки, задержки, номера, адреса Modbus) записываются числами
  Excel, а не текстом; десятичная запятая CSV понимается;
- строки-комментарии (первая ячейка начинается с ``#``) переносятся как есть;
- книга воспроизводима: время создания и время записей архива фиксированы, повторный запуск
  даёт побайтно тот же файл.

Запуск из корня репозитория::

    python examples/csv_to_xlsx.py examples/kns/io_list.csv examples/kns/io_list.xlsx \\
        --title "КНС-1. Перечень сигналов"
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import io
import re
import sys
import zipfile
from collections.abc import Sequence
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))  # запуск из клона без установки пакета

from iolistgen.columns import COLUMN_BY_NAME  # noqa: E402

# Колонки, которые пишутся числами Excel: целые и дробные
INT_COLUMNS = frozenset({"dev_no", "io_channel", "mb_slave", "mb_func", "mb_reg", "mb_bit"})
FLOAT_COLUMNS = frozenset({"range_min", "range_max", "sp_hh", "sp_h", "sp_l", "sp_ll", "delay_on_s", "delay_off_s"})

# Фиксированное время в свойствах книги и в записях архива: файл не меняется от запуска к запуску
FIXED_TIME = dt.datetime(2026, 1, 1, 0, 0, 0)
SHEET_TITLE = "Перечень"
MODIFIED_RE = re.compile(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)")
HEADER_FILL = PatternFill("solid", fgColor="DDE3EA")
COMMENT_FONT = Font(italic=True, color="6B7785")


def _number(column: str, text: str) -> object:
    """Значение ячейки: число для числовых колонок, иначе исходный текст."""
    value = text.strip()
    if not value or column not in INT_COLUMNS | FLOAT_COLUMNS:
        return value or None
    normalized = value.replace(",", ".")
    try:
        number = float(normalized)
    except ValueError:
        return value  # не число (например, «5 с» в колонке задержки) — оставить текстом
    if column in INT_COLUMNS and number.is_integer():
        return int(number)
    return int(number) if number.is_integer() and "." not in normalized else number


def read_csv(path: Path, delimiter: str) -> tuple[list[str], list[list[str]]]:
    """Заголовок и строки CSV (UTF-8, BOM допускается)."""
    with path.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream, delimiter=delimiter))
    if not rows:
        raise SystemExit(f"{path}: файл пуст")
    return rows[0], rows[1:]


def build_workbook(
    header: Sequence[str],
    rows: Sequence[Sequence[str]],
    *,
    title: str,
    subtitle: str,
    russian_headers: bool,
) -> Workbook:
    """Книга с одним листом «Перечень»: титул, строка заголовка, строки сигналов."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = SHEET_TITLE
    sheet.append([title])
    sheet.append([subtitle])
    sheet["A1"].font = Font(bold=True, size=13)
    sheet["A2"].font = Font(italic=True, color="56677D")

    titles = [COLUMN_BY_NAME[name].title if russian_headers and name in COLUMN_BY_NAME else name for name in header]
    sheet.append(titles)
    header_row = sheet.max_row
    for cell in sheet[header_row]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    widths = [len(text) for text in titles]
    for values in rows:
        if not any(value.strip() for value in values):
            continue
        comment = next((value for value in values if value.strip()), "").lstrip().startswith("#")
        cells = [
            (value.strip() or None) if comment else _number(header[index] if index < len(header) else "", value)
            for index, value in enumerate(values)
        ]
        sheet.append(cells)
        if comment:
            for cell in sheet[sheet.max_row]:
                cell.font = COMMENT_FONT
        for index, value in enumerate(values[: len(widths)]):
            widths[index] = max(widths[index], min(len(value), 60))

    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(6, width + 2)
    sheet.freeze_panes = f"B{header_row + 1}"
    last_column = get_column_letter(len(titles))
    sheet.auto_filter.ref = f"A{header_row}:{last_column}{sheet.max_row}"

    props = workbook.properties
    props.creator = "iolistgen examples"
    props.lastModifiedBy = "iolistgen examples"
    props.created = FIXED_TIME
    props.modified = FIXED_TIME
    props.title = title
    return workbook


def save_reproducible(workbook: Workbook, target: Path) -> None:
    """Сохранить книгу без меток текущего времени.

    openpyxl при записи ставит в ``docProps/core.xml`` время сохранения (``dcterms:modified``),
    а ZIP хранит время каждой записи архива: оба заменяются на ``FIXED_TIME``.
    """
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    result = io.BytesIO()
    stamp = FIXED_TIME.timetuple()[:6]
    modified = FIXED_TIME.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(result, "w", zipfile.ZIP_DEFLATED) as archive:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "docProps/core.xml":
                data = MODIFIED_RE.sub(lambda match: match.group(1) + modified + match.group(2), data)
            info = zipfile.ZipInfo(item.filename, date_time=stamp)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = item.external_attr
            archive.writestr(info, data)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(result.getvalue())


def _configure_streams() -> None:
    """Вывод в файл или канал — UTF-8 (в Windows по умолчанию кодировка консоли)."""
    for stream in (sys.stdout, sys.stderr):
        reconfigure = getattr(stream, "reconfigure", None)
        if reconfigure is not None and not stream.isatty():
            reconfigure(encoding="utf-8", errors="replace")


def main(argv: Sequence[str] | None = None) -> int:
    _configure_streams()
    parser = argparse.ArgumentParser(description="Перечень сигналов CSV → XLSX (титул, заголовки, числа).")
    parser.add_argument("source", type=Path, help="перечень CSV")
    parser.add_argument("target", type=Path, help="книга XLSX")
    parser.add_argument("--title", default="Перечень сигналов", help="первая строка титула")
    parser.add_argument("--delimiter", default=";", help="разделитель CSV (по умолчанию «;»)")
    parser.add_argument(
        "--headers",
        choices=("ru", "en"),
        default="ru",
        help="заголовки колонок: ru — русские названия схемы (по умолчанию), en — канонические имена",
    )
    args = parser.parse_args(argv)
    header, rows = read_csv(args.source, args.delimiter)
    subtitle = f"Собрано из {args.source.name} скриптом examples/csv_to_xlsx.py; схема — docs/schema.md"
    workbook = build_workbook(header, rows, title=args.title, subtitle=subtitle, russian_headers=args.headers == "ru")
    save_reproducible(workbook, args.target)
    print(f"{args.target}: {len(rows)} строк")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
